#!/usr/bin/env python3
"""
bbht_predicate500.py -- 보드에서 네 술어 x 500 워크로드를 자동으로 돌리고 엑셀로 저장

SW Common500 은 EQ 술어로 M = 1/4/16/64/256 x 공식 시드 100쌍 = 500 워크로드를
돌렸습니다. 이 프로그램은 같은 모양의 500 워크로드를 LT / GT / EQ / RANGE 에
각각 걸어 보드에서 돌리고, 결과를 SW 기준모델(Q1.22 bit-exact)의 기댓값과
워크로드마다 맞대어 .xlsx 로 남깁니다.

보드 쪽은 bbht_console 앱(2026-09-25 판 이상)이면 됩니다. 데이터셋은 UART 로
보내지 않고 보드가 GEN 으로 직접 만들며, 만든 배열이 기준모델과 같은지는 SUM 이
돌려주는 FNV-1a 해시로 확인합니다. 두 비트스트림 모두 됩니다.

  bbht_grover_upgrade (hardware_bram_K3H3_E4_M2)  NORMAL 과 CKPT(K3/H3-E4-M2) 를 같은 시드로
                                       짝지어 돌림 -> 2,000 x 2 = 4,000 실행
  bbht_grover_dram    (hardware_dram)  NORMAL 하나 -> 2,000 실행
  bbht_grover_nocheckpoint (hardware_bram/models/hardware_bram_nocheckpoint)
                                       NORMAL 하나 -> 2,000 실행. 체크포인트 없이
                                       연산기 네 벌(E4)만 켠 비교 기준

어느 쪽인지는 ID 응답의 platform= 으로 스스로 알아냅니다.

사용법 (보드 굽기와 앱 올리기를 마친 뒤)
  python3 software/host/bbht_predicate500.py --port /dev/ttyUSB1
  python3 software/host/bbht_predicate500.py --port COM4 --predicates EQ
  python3 software/host/bbht_predicate500.py --port mock          # 보드 없이 흐름 시험

  --out 을 안 주면 현재 폴더에 predicate500_<platform>_<시각>.xlsx 와 같은 이름의
  .csv 를 만듭니다. CSV 는 실행마다 한 줄씩 바로 써서, 중간에 끊겨도 거기까지는
  남습니다. --resume 으로 같은 CSV 를 주면 끝난 데이터셋은 건너뜁니다.

필요한 패키지: pyserial (보드), openpyxl (엑셀). numpy 는 필요 없습니다.
기댓값 파일은 software/experiments/predicate500_benchmark/expected/ 에 있습니다
(run_predicate500_golden.py 가 만듦).

비교 축 (워크로드마다)
  result_index, trial_count, L_BBHT   논리 궤적. 모드와 무관하게 기준모델과 같아야 함
  actual_iter                         물리 반복. bram NORMAL/CKPT, dram 세션 열, nocheckpoint NORMAL 열과 맞댐
  술어                                결과 인덱스의 값이 정말 술어를 만족하는가
시간 축 (기준모델에 없음. 보드에서만 나옴)
  cycle_count   가속기 사이클 (100 MHz)
  wall_us       보드 실시간 클럭으로 COMMAND 부터 DONE 을 본 순간까지 (보드 500런
                정본과 같은 구간). mock 에서는 지어낸 값입니다
"""

import argparse
import csv
import datetime as _dt
import hashlib
import os
import platform as _platform
import sys
import time
from collections import OrderedDict, defaultdict

_HERE = os.path.dirname(os.path.abspath(__file__))
_REPO = os.path.normpath(os.path.join(_HERE, os.pardir, os.pardir))
sys.path.insert(0, _HERE)

import bbht_cli  # noqa: E402  (같은 폴더. 트랜스포트와 파서를 같이 씁니다)

EXPECTED_DIR = os.path.join(_REPO, "software", "experiments", "predicate500_benchmark", "expected")
PRED_ORDER = ("LT", "GT", "EQ", "RANGE")
TARGETS = (1, 4, 16, 64, 256)
FW_MIN = "2026-09-25"

# 플랫폼 -> (갈래, 기본 모드 목록)
PLATFORMS = {
    "bbht_grover_upgrade": ("bram", ("normal", "ckpt")),
    "bbht_grover_dram":    ("dram", ("normal",)),
    "bbht_grover_nocheckpoint": ("nocheckpoint", ("normal",)),
}
# (갈래, 모드) -> 기댓값의 물리 반복 열
ITER_COLUMN = {
    ("bram", "normal"): "actual_iter_normal",
    ("bram", "ckpt"):   "actual_iter_k3h3",
    ("dram", "normal"): "actual_iter_dram_session",
    ("nocheckpoint", "normal"): "actual_iter_normal",
}

ROW_FIELDS = (
    "branch", "platform", "predicate", "threshold_a", "threshold_b", "target_count",
    "seed_index", "seed_j", "seed_meas", "mode", "outcome",
    "result_index", "result_value", "trial_count", "L_BBHT", "actual_iter",
    "cycle_count", "cycle_us", "wall_us",
    "exp_result_index", "exp_trial_count", "exp_L_BBHT", "exp_actual_iter",
    "match_result_index", "match_trial_count", "match_L_BBHT", "match_actual_iter",
    "predicate_ok", "all_match", "note",
)


# =====================================================================
# 기댓값
# =====================================================================
def load_expected(exp_dir):
    runs_path = os.path.join(exp_dir, "predicate500_expected.csv")
    ds_path = os.path.join(exp_dir, "predicate500_datasets.csv")
    for p in (runs_path, ds_path):
        if not os.path.exists(p):
            sys.exit("기댓값 파일이 없습니다: %s\n"
                     "software/experiments/predicate500_benchmark/run_predicate500_golden.py "
                     "를 먼저 돌리십시오." % p)
    with open(runs_path, encoding="utf-8") as f:
        runs = {(r["predicate"], int(r["target_count"]), int(r["seed_index"])): r
                for r in csv.DictReader(f)}
    with open(ds_path, encoding="utf-8") as f:
        datasets = {(r["predicate"], int(r["target_count"])): r for r in csv.DictReader(f)}
    digest = hashlib.sha256()
    for p in (runs_path, ds_path):
        with open(p, "rb") as f:
            digest.update(f.read())
    return runs, datasets, digest.hexdigest()


def seeds_from_expected(runs):
    """기댓값 CSV 에서 시드 로스터(순서대로)를 복원합니다."""
    seeds = {}
    for (p, m, s), r in runs.items():
        seeds.setdefault(s, (int(r["seed_j"], 16), int(r["seed_meas"], 16)))
    return [seeds[i] for i in sorted(seeds)]


def pred_ok(p, v, a, b):
    if p == "LT":
        return v < a
    if p == "GT":
        return v > a
    if p == "EQ":
        return v == a
    return a < v < b


# =====================================================================
# 보드 없이 흐름을 시험하는 mock
# =====================================================================
class GoldenMockTransport:
    """bbht_console 의 응답 모양을 흉내 냅니다. RUN 은 기댓값을 그대로 돌려주고
    사이클과 wall_us 는 지어냅니다 -- 파이프라인(파싱, 대조, 엑셀)을 시험하는
    용도이지 수치를 인용하는 용도가 아닙니다."""

    def __init__(self, runs, datasets, platform_name):
        self.runs, self.datasets = runs, datasets
        self.platform = platform_name
        self.branch = PLATFORMS.get(platform_name, ("bram", ()))[0]
        self.cfg = dict(MODE="EQ", A=0, B=0, COUNT=16384, CAP=100, SEEDJ=1, SEEDM=1,
                        AUTO=1, BURST=0, J=0, FAILLIM=4)
        self.cur = None          # (pred, M)
        self.loaded = False
        self.out = []

    @staticmethod
    def _num(s):
        neg = s.startswith("-")
        s = s[1:] if neg else s
        v = int(s, 16) if s.lower().startswith("0x") else int(s)
        return -v if neg else v

    def send(self, line):
        tok = line.upper().split()
        kv = dict(t.split("=", 1) for t in tok[1:] if "=" in t)
        cmd = tok[0] if tok else ""
        if cmd == "ID":
            self.out = ["STAT name=bbht_console csr_ver=0.9.8",
                        "STAT platform=%s fw=%s" % (self.platform, FW_MIN),
                        "STAT csr_base=0xe2020000 q_bits=14 n_entries=16384 fifo_depth=256",
                        "STAT accel_clk_hz=100000000 sram_base=0xe0000000 sram_last=0xe001ffff",
                        "STAT dataset_addr=0xe0004000 dataset_count=0 loaded=0", "OK"]
        elif cmd == "GEN":
            p, m = kv["PRED"], self._num(kv["TARGETS"])
            ds = self.datasets[(p, m)]
            self.cur, self.loaded = (p, m), False
            self.cfg.update(MODE=p, A=int(ds["threshold_a"]), B=int(ds["threshold_b"]))
            self.out = ["OK count=16384 targets=%d pred=%s a=%s b=%s"
                        % (m, p, ds["threshold_a"], ds["threshold_b"])]
        elif cmd == "SUM":
            ds = self.datasets[self.cur]
            self.out = ["STAT count=16384 fnv=%s hits=%d pred=%s a=%d b=%d loaded=%d"
                        % (ds["dataset_fnv1a"], self.cur[1], self.cfg["MODE"],
                           self.cfg["A"], self.cfg["B"], int(self.loaded)), "OK"]
        elif cmd == "LOAD":
            self.loaded = True
            self.out = ["OK loaded=16384 addr=0xe0004000"]
        elif cmd == "SET":
            for k, v in kv.items():
                self.cfg[k] = v if k == "MODE" else self._num(v)
            self.out = ["OK"]
        elif cmd == "RUN":
            p, m = self.cur
            seeds = {(int(r["seed_j"], 16), int(r["seed_meas"], 16)): s
                     for (pp, mm, s), r in self.runs.items() if pp == p and mm == m}
            s = seeds[(self.cfg["SEEDJ"], self.cfg["SEEDM"])]
            r = self.runs[(p, m, s)]
            mode = "ckpt" if self.cfg["BURST"] else "normal"
            it = int(r[ITER_COLUMN[(self.branch, mode)]])
            cyc = 700 * it + 900 * int(r["trial_count"])
            self.out = ["HIT idx=%s val=%s trials=%s l=%s iters=%d cyc=%d us=%d wall_us=%d"
                        % (r["result_index"], r["result_value"], r["trial_count"],
                           r["L_BBHT"], it, cyc, cyc // 100, cyc // 100 + 12), "OK"]
        else:
            self.out = ["ERR UNKNOWN_CMD"]

    def read_lines(self):
        return self.out

    def close(self):
        pass


# =====================================================================
# 보드와 주고받기
# =====================================================================
def terminator(lines):
    return lines[-1] if lines else "(무응답)"


def expect_ok(board, line):
    lines = board.cmd(line)
    if not lines or not lines[-1].startswith("OK"):
        raise RuntimeError("명령 실패: %s -> %s" % (line, terminator(lines)))
    return lines


def board_identity(board):
    ident = {}
    raw = board.cmd("ID")
    for line in raw:
        tag, kv = bbht_cli.parse_kv_line(line)
        if tag == "STAT":
            ident.update(kv)
    return ident, raw


def run_once(board, seed_j, seed_meas, burst):
    expect_ok(board, "SET BURST=%d SEEDJ=0x%08X SEEDM=0x%08X" % (burst, seed_j, seed_meas))
    lines = board.cmd("RUN")
    if not lines or not lines[-1].startswith("OK"):
        return "ERR", {}, terminator(lines)
    for line in lines:
        tag, kv = bbht_cli.parse_kv_line(line)
        if tag in ("HIT", "MISS"):
            return tag, kv, ""
    return "ERR", {}, "HIT/MISS 줄이 없음"


def compare(row, exp, branch, mode):
    a, b = int(exp["threshold_a"]), int(exp["threshold_b"])
    col = ITER_COLUMN.get((branch, mode))
    success = exp["success"] == "1"
    row["exp_result_index"] = exp["result_index"] if success else ""
    row["exp_trial_count"] = int(exp["trial_count"])
    row["exp_L_BBHT"] = int(exp["L_BBHT"])
    row["exp_actual_iter"] = int(exp[col]) if col else ""
    hit = row["outcome"] == "HIT"
    row["match_result_index"] = int(
        (hit and success and int(row["result_index"]) == int(exp["result_index"]))
        or (not hit and not success and row["outcome"] == "MISS"))
    row["match_trial_count"] = int(row["trial_count"] != "" and
                                   int(row["trial_count"]) == row["exp_trial_count"])
    row["match_L_BBHT"] = int(row["L_BBHT"] != "" and int(row["L_BBHT"]) == row["exp_L_BBHT"])
    row["match_actual_iter"] = int(col is not None and row["actual_iter"] != "" and
                                   int(row["actual_iter"]) == row["exp_actual_iter"])
    row["predicate_ok"] = int((not hit) or pred_ok(row["predicate"], int(row["result_value"]), a, b))
    row["all_match"] = int(all(row[k] for k in ("match_result_index", "match_trial_count",
                                                "match_L_BBHT", "match_actual_iter",
                                                "predicate_ok")))


# =====================================================================
# 요약과 엑셀
# =====================================================================
def summarize(rows, modes):
    """술어 x 모드 요약, 술어 x M x 모드 요약."""
    by_pm = OrderedDict()
    by_pmm = OrderedDict()
    for p in PRED_ORDER:
        for mode in modes:
            by_pm[(p, mode)] = []
            for m in TARGETS:
                by_pmm[(p, m, mode)] = []
    for r in rows:
        key = (r["predicate"], r["mode"])
        if key in by_pm:
            by_pm[key].append(r)
            by_pmm[(r["predicate"], int(r["target_count"]), r["mode"])].append(r)

    def agg(items):
        n = len(items)
        s = lambda k: sum(int(x[k]) for x in items if x[k] != "")
        cyc = s("cycle_count")
        wall = s("wall_us")
        return OrderedDict([
            ("runs", n),
            ("hit", sum(1 for x in items if x["outcome"] == "HIT")),
            ("match_result_index", s("match_result_index")),
            ("match_trial_count", s("match_trial_count")),
            ("match_L_BBHT", s("match_L_BBHT")),
            ("match_actual_iter", s("match_actual_iter")),
            ("predicate_ok", s("predicate_ok")),
            ("all_match", s("all_match")),
            ("sum_actual_iter", s("actual_iter")),
            ("sum_cycle_count", cyc),
            ("sum_wall_us", wall),
            ("mean_wall_us", round(wall / n, 1) if n else 0),
        ])

    return ([dict(predicate=p, mode=mode, **agg(v)) for (p, mode), v in by_pm.items() if v],
            [dict(predicate=p, target_count=m, mode=mode, **agg(v))
             for (p, m, mode), v in by_pmm.items() if v])


def write_xlsx(path, rows, ds_rows, summary, summary_m, env, modes, is_mock):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl 이 없어 엑셀을 못 만들었습니다 (pip install openpyxl). CSV 는 남았습니다.")
        return False

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDE3EA")
    bad_fill = PatternFill("solid", fgColor="F8D0D0")
    ok_fill = PatternFill("solid", fgColor="D8EED8")
    warn_fill = PatternFill("solid", fgColor="FFF2B3")

    def sheet(ws, header, data, widths=None, bad_col=None):
        ws.append(list(header))
        for c in ws[1]:
            c.font, c.fill = bold, head_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for d in data:
            ws.append([d.get(h, "") if isinstance(d, dict) else d[i] for i, h in enumerate(header)])
            if bad_col is not None:
                v = d.get(bad_col) if isinstance(d, dict) else None
                if v in (0, "0", False):
                    for c in ws[ws.max_row]:
                        c.fill = bad_fill
        ws.freeze_panes = "A2"
        for i, h in enumerate(header, start=1):
            w = (widths or {}).get(h, max(10, min(28, len(str(h)) + 2)))
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()

    # 1. 요약
    ws = wb.active
    ws.title = "요약"
    if is_mock:
        ws.append(["MOCK 실행입니다. 보드 없이 파이프라인만 시험한 결과이므로 수치를 인용하지 마십시오."])
        ws["A1"].font = Font(bold=True, color="C00000")
        ws.append([])
    ws.append(["플랫폼", env.get("platform", ""), "갈래", env.get("branch", ""),
               "실행 시각", env.get("started", "")])
    ws.append([])
    header = ["predicate", "mode", "runs", "hit", "match_result_index", "match_trial_count",
              "match_L_BBHT", "match_actual_iter", "predicate_ok", "all_match",
              "sum_actual_iter", "sum_cycle_count", "sum_wall_us", "mean_wall_us"]
    start = ws.max_row + 1
    ws.append(header)
    for c in ws[start]:
        c.font, c.fill = bold, head_fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for s in summary:
        ws.append([s[h] for h in header])
        fill = ok_fill if s["all_match"] == s["runs"] else bad_fill
        ws.cell(row=ws.max_row, column=header.index("all_match") + 1).fill = fill
    for i, h in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(11, len(h) + 2)

    # 모드 짝 배수 (bram 만)
    if "normal" in modes and "ckpt" in modes:
        ws.append([])
        ws.append(["술어", "NORMAL/CKPT 사이클 배수", "NORMAL/CKPT 실경과 배수",
                   "NORMAL 물리반복", "CKPT 물리반복"])
        for c in ws[ws.max_row]:
            c.font, c.fill = bold, head_fill
        for p in PRED_ORDER:
            n = next((s for s in summary if s["predicate"] == p and s["mode"] == "normal"), None)
            k = next((s for s in summary if s["predicate"] == p and s["mode"] == "ckpt"), None)
            if not n or not k:
                continue
            ws.append([p,
                       round(n["sum_cycle_count"] / k["sum_cycle_count"], 4) if k["sum_cycle_count"] else "",
                       round(n["sum_wall_us"] / k["sum_wall_us"], 4) if k["sum_wall_us"] else "",
                       n["sum_actual_iter"], k["sum_actual_iter"]])

    ws.append([])
    notes = [
        "비교 기준: software/experiments/predicate500_benchmark/expected/ (Q1.22 bit-exact 기준모델)",
        "all_match = result_index, trial_count, L_BBHT, actual_iter 가 기댓값과 같고 결과 값이 술어를 만족",
        "cycle_count 는 가속기 100 MHz 사이클, wall_us 는 보드 실시간 클럭(1 MHz)으로 COMMAND~DONE 을 잰 값",
        "고전 선형 스캔과 비교하지 마십시오. 에뮬레이터는 고전 스캔보다 느립니다 (CLAUDE.md 2절)",
    ]
    if env.get("branch") == "nocheckpoint":
        notes.append("nocheckpoint 는 체크포인트·정책 엔진·측정 최적화(M1/M2) 없이 연산기 네 벌(E4)만 켠 "
                     "비교 기준 판입니다. actual_iter 는 기준모델 NORMAL 열과 맞댑니다")
    if env.get("branch") == "dram":
        notes.append("dram 의 actual_iter 는 데이터셋 하나를 적재한 뒤 시드를 차례로 돌며 DRAM 표를 "
                     "이어 쓴 값(기준모델 actual_iter_dram_session)과 맞댑니다")
    for n in notes:
        ws.append([n])

    # 2. M별 요약
    ws = wb.create_sheet("M별 요약")
    sheet(ws, ["predicate", "target_count", "mode", "runs", "hit", "all_match",
               "match_result_index", "match_trial_count", "match_L_BBHT", "match_actual_iter",
               "predicate_ok", "sum_actual_iter", "sum_cycle_count", "sum_wall_us", "mean_wall_us"],
          summary_m)

    # 3. 실행별
    ws = wb.create_sheet("실행별")
    sheet(ws, ROW_FIELDS, rows, widths={"note": 30, "platform": 20}, bad_col="all_match")

    # 4. 데이터셋
    ws = wb.create_sheet("데이터셋")
    sheet(ws, ["predicate", "target_count", "threshold_a", "threshold_b", "gen_command",
               "exp_fnv1a", "board_fnv1a", "board_hits", "fnv_match", "hits_match"],
          ds_rows, widths={"gen_command": 70}, bad_col="fnv_match")

    # 5. 환경
    ws = wb.create_sheet("환경")
    ws.append(["항목", "값"])
    for c in ws[1]:
        c.font, c.fill = bold, head_fill
    for k, v in env.items():
        if k == "id_lines":
            continue
        ws.append([k, str(v)])
    ws.append([])
    ws.append(["ID 응답"])
    for line in env.get("id_lines", []):
        ws.append(["", line])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100

    if any(r["all_match"] in (0, "0") for r in rows):
        wb["요약"].sheet_properties.tabColor = "C00000"
    wb.save(path)
    return True


# =====================================================================
# 본체
# =====================================================================
def parse_list(text, allowed, conv=str):
    items = [conv(x.strip().upper() if conv is str else x.strip())
             for x in text.split(",") if x.strip()]
    for x in items:
        if x not in allowed:
            sys.exit("허용되지 않는 값 %r. 가능한 값: %s" % (x, ", ".join(map(str, allowed))))
    return items


def emit_script(path, commands):
    """SoC RTL 시뮬용 bbht_console 스크립트 헤더를 씁니다."""
    with open(path, "w", encoding="utf-8") as f:
        f.write("/* bbht_predicate500.py --emit-script 가 만든 콘솔 스크립트.\n"
                " * RTL 시뮬 빌드에서 BBHT_CONSOLE_SCRIPT_FILE 로 넘깁니다.\n"
                " * 호스트가 보드에 보낼 명령과 한 줄도 다르지 않아야 트랜스크립트를\n"
                " * --port replay:<qtsim.log> 로 되먹일 수 있습니다. */\n")
        f.write("#define BBHT_CONSOLE_SCRIPT_LINES \\\n")
        for c in commands:
            f.write('    "%s", \\\n' % c)
        f.write('    "QUIT"\n')
    print("스크립트 %d 줄 -> %s" % (len(commands) + 1, path))


class RecordingBoard:
    """--emit-script 용. 보낸 명령만 모으고 응답은 mock 에서 받습니다."""

    def __init__(self, inner):
        self.inner = inner
        self.sent = []

    def cmd(self, line):
        self.sent.append(line)
        return self.inner.cmd(line)


def main():
    ap = argparse.ArgumentParser(
        description="보드에서 네 술어 x 500 워크로드를 자동으로 돌려 SW 기준모델과 맞대고 엑셀로 저장")
    ap.add_argument("--port", required=True,
                    help="시리얼 장치(/dev/ttyUSB1, COM4), 'mock' 또는 'replay:<트랜스크립트>'")
    ap.add_argument("--baud", type=int, default=115200)
    ap.add_argument("--timeout", type=float, default=60.0, help="명령 하나의 응답 대기 상한(초)")
    ap.add_argument("--predicates", default="LT,GT,EQ,RANGE")
    ap.add_argument("--targets", default="1,4,16,64,256")
    ap.add_argument("--seeds", type=int, default=100, help="앞에서부터 몇 쌍 (정본은 100)")
    ap.add_argument("--modes", default="auto",
                    help="auto(플랫폼으로 결정) 또는 normal / normal,ckpt")
    ap.add_argument("--out", help="엑셀 경로. 같은 이름의 .csv 도 같이 만듭니다")
    ap.add_argument("--resume", action="store_true",
                    help="--out 의 CSV 가 있으면 끝난 데이터셋은 건너뜁니다")
    ap.add_argument("--expected-dir", default=EXPECTED_DIR)
    ap.add_argument("--mock-platform", default="bbht_grover_upgrade", choices=sorted(PLATFORMS),
                    help="--port mock 일 때 흉내 낼 플랫폼")
    ap.add_argument("--emit-script", metavar="HEADER",
                    help="보드 대신 mock 으로 명령 순서만 뽑아 SoC 시뮬용 스크립트 헤더로 씁니다")
    ap.add_argument("--log", help="주고받은 줄 전부를 기록할 파일")
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args()

    runs, datasets, exp_sha = load_expected(args.expected_dir)
    predicates = parse_list(args.predicates, PRED_ORDER)
    targets = parse_list(args.targets, TARGETS, conv=int)
    seeds = seeds_from_expected(runs)[: args.seeds]

    is_mock = args.port == "mock" or bool(args.emit_script)
    if is_mock:
        transport = GoldenMockTransport(runs, datasets, args.mock_platform)
    elif args.port.startswith("replay:"):
        transport = bbht_cli.ReplayTransport(args.port.split(":", 1)[1])
    else:
        transport = bbht_cli.SerialTransport(args.port, args.baud, args.timeout)
    board = bbht_cli.Board(transport)
    if args.emit_script:
        board = RecordingBoard(board)

    logf = open(args.log, "w", encoding="utf-8") if args.log else None
    if logf:
        inner_cmd = board.cmd

        def logged(line):
            out = inner_cmd(line)
            logf.write("> %s\n" % line)
            for l in out:
                logf.write("%s\n" % l)
            logf.flush()
            return out
        board.cmd = logged

    started = _dt.datetime.now()
    try:
        # ---------- 1. 보드 확인 ----------
        ident, id_lines = board_identity(board)
        plat = str(ident.get("platform", ""))
        fw = str(ident.get("fw", ""))
        if not plat:
            sys.exit("ID 응답에 platform= 이 없습니다. bbht_console 을 2026-09-25 판 이상으로 "
                     "다시 올리십시오 (GEN PRED / SUM / wall_us 가 필요합니다).\n  응답: %s"
                     % " | ".join(id_lines))
        if fw < FW_MIN:
            sys.exit("콘솔 펌웨어가 %s 입니다. %s 이상이 필요합니다." % (fw, FW_MIN))
        if int(ident.get("n_entries", 0)) != 16384:
            sys.exit("n_entries=%s. Q14(16384) 비트스트림이 아닙니다." % ident.get("n_entries"))
        branch, default_modes = PLATFORMS.get(plat, ("bram", ("normal", "ckpt")))
        if args.modes == "auto":
            modes = list(default_modes)
        else:
            modes = [x.strip().lower() for x in args.modes.split(",") if x.strip()]
            if not modes or any(x not in ("normal", "ckpt") for x in modes):
                sys.exit("--modes 는 auto, normal, normal,ckpt 중 하나입니다")
        if "ckpt" in modes and (branch, "ckpt") not in ITER_COLUMN:
            sys.exit("%s 갈래에는 체크포인트 모드가 없습니다. --modes normal 로 돌리십시오." % branch)

        stamp = started.strftime("%Y%m%d_%H%M%S")
        out_xlsx = args.out or os.path.join(os.getcwd(), "predicate500_%s_%s.xlsx" % (plat, stamp))
        if not out_xlsx.lower().endswith(".xlsx"):
            out_xlsx += ".xlsx"
        out_csv = out_xlsx[:-5] + ".csv"

        total = len(predicates) * len(targets) * len(seeds) * len(modes)
        if not args.quiet:
            print("보드: platform=%s (%s 갈래) fw=%s" % (plat, branch, fw))
            print("계획: 술어 %s x M %s x 시드 %d x 모드 %s = %d 실행"
                  % ("/".join(predicates), "/".join(map(str, targets)), len(seeds),
                     "/".join(modes), total))
            if not args.emit_script:
                print("결과: %s (+ .csv)" % out_xlsx)

        # ---------- 이어 하기 ----------
        rows, ds_rows, done_ds = [], [], set()
        if args.resume and os.path.exists(out_csv):
            with open(out_csv, encoding="utf-8") as f:
                prev = list(csv.DictReader(f))
            count = defaultdict(int)
            for r in prev:
                count[(r["predicate"], int(r["target_count"]))] += 1
            for key, n in count.items():
                if n == len(seeds) * len(modes):
                    done_ds.add(key)
            rows = [r for r in prev if (r["predicate"], int(r["target_count"])) in done_ds]
            print("이어 하기: 끝난 데이터셋 %d 개, 실행 %d 줄을 살립니다" % (len(done_ds), len(rows)))

        csvf = None
        if not args.emit_script:
            csvf = open(out_csv, "w", newline="", encoding="utf-8")
            writer = csv.DictWriter(csvf, fieldnames=ROW_FIELDS, lineterminator="\n")
            writer.writeheader()
            writer.writerows(rows)
            csvf.flush()

        # ---------- 2. 술어 x M ----------
        done = len(rows)
        t0 = time.time()
        for p in predicates:
            for m in targets:
                ds = datasets[(p, m)]
                a, b = int(ds["threshold_a"]), int(ds["threshold_b"])
                ds_row = OrderedDict(predicate=p, target_count=m, threshold_a=a, threshold_b=b,
                                     gen_command=ds["gen_command"], exp_fnv1a=ds["dataset_fnv1a"],
                                     board_fnv1a="", board_hits="", fnv_match=0, hits_match=0)
                if (p, m) in done_ds:
                    ds_row["board_fnv1a"] = "(이어 하기로 건너뜀)"
                    ds_rows.append(ds_row)
                    continue

                expect_ok(board, ds["gen_command"])
                sums = {}
                for line in expect_ok(board, "SUM"):
                    tag, kv = bbht_cli.parse_kv_line(line)
                    if tag == "STAT":
                        sums.update(kv)
                ds_row["board_fnv1a"] = "0x%08x" % sums.get("fnv", 0) if "fnv" in sums else ""
                ds_row["board_hits"] = sums.get("hits", "")
                ds_row["fnv_match"] = int(sums.get("fnv") == int(ds["dataset_fnv1a"], 16))
                ds_row["hits_match"] = int(sums.get("hits") == m)
                ds_rows.append(ds_row)
                if not (ds_row["fnv_match"] and ds_row["hits_match"]):
                    print("  %s M=%d: 보드 데이터셋이 기준모델과 다릅니다 (fnv %s / 기댓값 %s, "
                          "hits %s). 이 데이터셋은 건너뜁니다."
                          % (p, m, ds_row["board_fnv1a"], ds["dataset_fnv1a"], ds_row["board_hits"]))
                    continue

                expect_ok(board, "LOAD")
                expect_ok(board, "SET MODE=%s A=%d B=%d COUNT=16384 CAP=100 AUTO=1 J=0" % (p, a, b))

                n_ok = 0
                for s, (sj, sm) in enumerate(seeds):
                    exp = runs[(p, m, s)]
                    for mode in modes:
                        tag, kv, note = run_once(board, sj, sm, 1 if mode == "ckpt" else 0)
                        row = OrderedDict((k, "") for k in ROW_FIELDS)
                        row.update(branch=branch, platform=plat, predicate=p, threshold_a=a,
                                   threshold_b=b, target_count=m, seed_index=s,
                                   seed_j="0x%08x" % sj, seed_meas="0x%08x" % sm,
                                   mode=mode, outcome=tag, note=note)
                        if tag in ("HIT", "MISS"):
                            row.update(result_index=kv.get("idx", "") if tag == "HIT" else "",
                                       result_value=kv.get("val", "") if tag == "HIT" else "",
                                       trial_count=kv.get("trials", ""),
                                       L_BBHT=kv.get("l", ""),
                                       actual_iter=kv.get("iters", ""),
                                       cycle_count=kv.get("cyc", ""),
                                       cycle_us=kv.get("us", ""),
                                       wall_us=kv.get("wall_us", ""))
                            if tag == "MISS":
                                row["note"] = "reason=%s" % kv.get("reason", "")
                            compare(row, exp, branch, mode)
                        else:
                            for k in ("match_result_index", "match_trial_count", "match_L_BBHT",
                                      "match_actual_iter", "predicate_ok", "all_match"):
                                row[k] = 0
                        n_ok += int(row["all_match"])
                        rows.append(row)
                        if csvf:
                            writer.writerow(row)
                        done += 1
                if csvf:
                    csvf.flush()
                if not args.quiet:
                    el = time.time() - t0
                    print("  %-5s M=%-3d  일치 %4d/%-4d  (누적 %d/%d, %.0f초)"
                          % (p, m, n_ok, len(seeds) * len(modes), done, total, el))
    finally:
        if logf:
            logf.close()
        transport.close()

    if args.emit_script:
        emit_script(args.emit_script, board.sent)
        return 0

    csvf.close()
    summary, summary_m = summarize(rows, modes)
    env = OrderedDict(
        started=started.isoformat(timespec="seconds"),
        finished=_dt.datetime.now().isoformat(timespec="seconds"),
        port="mock" if is_mock else args.port,
        platform=plat, branch=branch, firmware=fw, modes="/".join(modes),
        predicates="/".join(predicates), targets="/".join(map(str, targets)), seeds=len(seeds),
        expected_dir=os.path.relpath(args.expected_dir, _REPO),
        expected_sha256=exp_sha,
        host=_platform.node(), python=_platform.python_version(),
        command=" ".join(sys.argv),
        id_lines=id_lines,
    )
    ok = write_xlsx(out_xlsx, rows, ds_rows, summary, summary_m, env, modes, is_mock)

    # 터미널 요약
    print("")
    print("%-6s %-7s %6s %10s %14s %14s" % ("술어", "모드", "실행", "전부일치", "사이클 합", "wall_us 합"))
    for s in summary:
        print("%-6s %-7s %6d %10s %14d %14d"
              % (s["predicate"], s["mode"], s["runs"], "%d/%d" % (s["all_match"], s["runs"]),
                 s["sum_cycle_count"], s["sum_wall_us"]))
    bad = sum(1 for r in rows if str(r["all_match"]) != "1")
    bad_ds = sum(1 for d in ds_rows if d["fnv_match"] == 0 and "건너뜀" not in str(d["board_fnv1a"]))
    print("")
    print("판정: %s  (불일치 실행 %d, 데이터셋 불일치 %d)%s"
          % ("PASS" if bad == 0 and bad_ds == 0 and rows else "FAIL", bad, bad_ds,
             "  [MOCK - 수치 인용 금지]" if is_mock else ""))
    print("CSV : %s" % out_csv)
    if ok:
        print("엑셀: %s" % out_xlsx)
    return 0 if bad == 0 and bad_ds == 0 and rows else 1


if __name__ == "__main__":
    sys.exit(main())
